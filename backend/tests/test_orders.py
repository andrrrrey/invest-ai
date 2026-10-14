"""Пункт 13 КП: проект приказа (ЛНА) по согласованному контракту."""

import io
import json

from docx import Document

from _helpers import client, mk_user, auth, mk_contract, setup_route_roles, get_project
from app import settings_store
from app.services import ai_service, mattermost_service, order_service


def _approved_contract(owner):
    return mk_contract(owner, "leadership", "approved", curator="Кураторов К.", customerName="Заказчиков И.",
                       startDate="2026-11-01", deadline="2027-03-31",
                       milestones=[{"name": "Выручка филиала", "deadline": "2027-01-31", "criteria": "Выручка ≥ 10 млн ₽",
                                    "baseline": "0", "target": "10 млн ₽", "dataSource": "1С"}])


def _docx_text(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


def test_template_order_and_route(monkeypatch):
    monkeypatch.setattr(mattermost_service, "is_configured", lambda: False)
    monkeypatch.setattr(settings_store, "is_ai_enabled", lambda: False)
    roles = setup_route_roles()
    owner = mk_user("owner")
    pid = _approved_contract(owner)

    r = client.post(f"/api/v1/projects/{pid}/order", headers=auth(owner))
    assert r.status_code == 200, r.text
    state = r.json()
    assert state["order"]["source"] == "template"
    assert [s["role"] for s in state["route"]["steps"]] == ["cb", "lawyer", "ceo"]
    att_id = state["attachment"]["id"]

    d = client.get(f"/api/v1/attachments/{att_id}/download", headers=auth(owner))
    text = _docx_text(d.content)
    assert "НЕПУБЛИЧНОЕ АКЦИОНЕРНОЕ ОБЩЕСТВО" in text and "ПРИКАЗЫВАЮ" in text
    assert "Выручка ≥ 10 млн ₽" in text and "А.М. Пелевин" in text
    assert "Контроль за исполнением настоящего приказа оставляю за собой." in text

    # Повторно сформировать нельзя, пока идёт согласование.
    assert client.post(f"/api/v1/projects/{pid}/order", headers=auth(owner)).status_code == 409

    steps = {s["role"]: s["id"] for s in state["route"]["steps"]}
    # Юрист (роль owner) видит карточку и скачивает приказ, когда очередь дошла.
    client.post(f"/api/v1/projects/{pid}/route/steps/{steps['cb']}/decision", headers=auth(roles["cb"]), json={"decision": "approve"})
    assert client.get(f"/api/v1/attachments/{att_id}/download", headers=auth(roles["lawyer"])).status_code == 200
    r = client.post(f"/api/v1/projects/{pid}/route/steps/{steps['lawyer']}/decision", headers=auth(roles["lawyer"]),
                    json={"decision": "rework", "comment": "Уточнить срок ознакомления"})
    assert r.status_code == 200
    # Контракт остаётся согласованным — на доработку уходит только приказ.
    assert get_project(pid).status == "approved"

    # Исправленный файл → новый цикл согласования.
    up = client.post(f"/api/v1/projects/{pid}/order/upload", headers=auth(owner),
                     files={"file": ("Приказ v2.docx", d.content, "application/octet-stream")})
    assert up.status_code == 200, up.text
    st = up.json()["route"]
    assert st["cycle"] == 2 and st["steps"][0]["status"] == "active"
    ids = {s["role"]: s["id"] for s in st["steps"]}
    for role in ("cb", "lawyer", "ceo"):
        assert client.post(f"/api/v1/projects/{pid}/route/steps/{ids[role]}/decision", headers=auth(roles[role]),
                           json={"decision": "approve"}).status_code == 200
    final = client.get(f"/api/v1/projects/{pid}/order", headers=auth(owner)).json()
    assert final["route"]["state"] == "approved" and final["order"].get("approved_at")


def test_ai_order_uses_skill(monkeypatch):
    monkeypatch.setattr(settings_store, "is_ai_enabled", lambda: True)
    seen = {}

    def fake_chat(prompt, max_tokens=700, **kw):
        seen["prompt"], seen["max_tokens"] = prompt, max_tokens
        return json.dumps({
            "title": "Об утверждении плана работ и постановке задач Иванова И.И.",
            "preamble": "В целях конкретизации трудовой функции..., на основании ст. 15, 274 Трудового кодекса РФ,",
            "points": ["Утвердить план работ.", "Контроль за исполнением настоящего приказа оставляю за собой."],
            "appendix_title": "План работ", "plan_rows": [{"result": "R", "deadline": "31.01.2027",
                                                          "criterion": "C", "resources": "Res", "confirmation": "Rep"}],
            "review_notes": ["Проверить сроки"],
        })
    monkeypatch.setattr(ai_service, "_chat", fake_chat)
    owner = mk_user("owner")
    project = get_project(_approved_contract(owner))
    order, source = order_service.build_order(None, project)
    assert source == "ai" and order["appendix_rows"][0][1] == "R"
    assert "Постановка задач через ЛНА" in seen["prompt"] and seen["max_tokens"] >= 4000
    assert "ПРИКАЗЫВАЮ" in _docx_text(order_service.build_order_docx(order))


def test_order_requires_approved_contract(monkeypatch):
    monkeypatch.setattr(mattermost_service, "is_configured", lambda: False)
    setup_route_roles()
    owner = mk_user("owner")
    pid = mk_contract(owner, "smart", "draft")
    assert client.post(f"/api/v1/projects/{pid}/order", headers=auth(owner)).status_code == 400
    stranger = mk_user("owner")
    assert client.get(f"/api/v1/projects/{pid}/order", headers=auth(stranger)).status_code == 403
