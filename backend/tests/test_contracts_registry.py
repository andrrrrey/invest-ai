"""Пункт 11 КП: отчёт «Реестр контрактов»."""

import io

from openpyxl import load_workbook

from _helpers import client, mk_user, auth, mk_contract


def test_registry_excel_with_filters():
    cfo = mk_user("cfo")
    owner = mk_user("owner")
    smart = mk_contract(owner, "smart", "approved", navRate=10.0, customerName="Заказчиков",
                        milestones=[{"name": "Э1", "rewardRub": 1000, "coins": 100, "status": "paid"},
                                    {"name": "Э2", "rewardRub": 500, "coins": 50}])
    lead = mk_contract(owner, "leadership", "draft", leadershipTotals={"grand": 900000})

    r = client.post("/api/v1/export/contracts-registry", headers=auth(cfo),
                    json={"format": "excel", "owner_id": owner})
    assert r.status_code == 200, r.text
    assert "filename*=UTF-8''" in r.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(r.content)).active
    header = [c.value for c in ws[4]]
    assert header[:5] == ["№", "Контракт", "Тип", "Статус", "Этап согласования"]
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=5, max_row=6)]
    by_type = {row[2]: row for row in rows}
    assert by_type["Смарт-контракт"][9] == 1500 and by_type["Смарт-контракт"][10] == 150
    assert by_type["Смарт-контракт"][12] == "1/2"
    assert by_type["Лидерский контракт"][9] == 900000

    r = client.post("/api/v1/export/contracts-registry", headers=auth(cfo),
                    json={"format": "excel", "owner_id": owner, "sc_type": "leadership", "statuses": ["draft"]})
    ws = load_workbook(io.BytesIO(r.content)).active
    names = [ws.cell(row=i, column=3).value for i in range(5, ws.max_row)]
    assert names == ["Лидерский контракт"]


def test_registry_pdf_and_owner_blocked():
    cfo = mk_user("cfo")
    r = client.post("/api/v1/export/contracts-registry", headers=auth(cfo), json={"format": "pdf"})
    assert r.status_code == 200 and r.content[:4] == b"%PDF"
    owner = mk_user("owner")
    assert client.post("/api/v1/export/contracts-registry", headers=auth(owner), json={}).status_code == 403
