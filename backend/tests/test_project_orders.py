"""Пункт 14 КП: проект приказа по согласованному инвест-проекту."""

import io

from docx import Document

from _helpers import client, mk_user, auth, setup_route_roles
from app import settings_store
from app.database import SessionLocal
from app.models.project import Project
from app.models.tranche import Tranche
from app.services import mattermost_service


def _approved_investment(owner_id):
    s = SessionLocal()
    try:
        p = Project(project_type="investment", name="Новый склад", business_unit="Логистика",
                    owner="Складов С.", user_id=owner_id, status="approved",
                    financial_model={"initialInvestment": -5000000})
        s.add(p)
        s.flush()
        s.add_all([
            Tranche(project_id=p.id, amount=3000000, planned_date="2026-12-01", status="approved",
                    description="Аренда и ремонт", order_index=0),
            Tranche(project_id=p.id, amount=2000000, planned_date="2027-03-01", status="requested",
                    description="Оборудование", order_index=1),
        ])
        s.commit()
        return p.id
    finally:
        s.close()


def test_investment_project_order(monkeypatch):
    monkeypatch.setattr(mattermost_service, "is_configured", lambda: False)
    monkeypatch.setattr(settings_store, "is_ai_enabled", lambda: False)
    setup_route_roles()
    owner = mk_user("owner")
    pid = _approved_investment(owner)
    r = client.post(f"/api/v1/projects/{pid}/order", headers=auth(owner))
    assert r.status_code == 200, r.text
    state = r.json()
    assert [s["role"] for s in state["route"]["steps"]] == ["cb", "lawyer", "ceo"]
    content = client.get(f"/api/v1/attachments/{state['attachment']['id']}/download", headers=auth(owner)).content
    doc = Document(io.BytesIO(content))
    text = "\n".join(p.text for p in doc.paragraphs)
    cells = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "О выделении денежных средств на реализацию инвестиционного проекта «Новый склад»" in text
    assert "5 000 000" in text and "Складов С." in text
    assert "3 000 000" in cells and "Оборудование" in cells and "01.03.2027" in cells
