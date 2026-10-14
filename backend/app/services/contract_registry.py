"""Отчёт «Реестр контрактов»: выборка с фильтрами и выгрузка в Excel/PDF."""

from __future__ import annotations

import datetime as dt
import html
import io
from typing import List, Optional

from sqlalchemy import or_
from sqlalchemy.orm import Session

from ..models.project import Project
from ..models.user import User
from . import route_service
from .contract_utils import SC_TYPE_LABELS, sc_type_of

STATUS_LABELS = {
    "draft": "Черновик",
    "pending_approval": "На согласовании",
    "approved": "Согласован",
    "rejected": "Отклонён",
    "rework_needed": "На доработке",
}

COLUMNS = [
    ("num", "№"), ("name", "Контракт"), ("type", "Тип"), ("status", "Статус"),
    ("stage", "Этап согласования"), ("owner", "Ответственный"), ("customer", "Заказчик"),
    ("curator", "Куратор"), ("business_unit", "Бизнес-юнит"), ("reward_rub", "Вознаграждение, ₽"),
    ("coins", "Coin"), ("rate", "Курс Coin, ₽"), ("milestones", "Этапы (выполнено/всего)"),
    ("created", "Создан"), ("deadline", "Дедлайн"),
]


def _date(value) -> Optional[dt.date]:
    if not value:
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    try:
        return dt.date.fromisoformat(str(value)[:10])
    except ValueError:
        return None


def _fmt_date(value) -> str:
    d = _date(value)
    return d.strftime("%d.%m.%Y") if d else ""


def _num(value) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def collect(
    db: Session,
    user: User,
    *,
    sc_type: Optional[str] = None,
    statuses: Optional[List[str]] = None,
    date_from: Optional[dt.date] = None,
    date_to: Optional[dt.date] = None,
    owner_id: Optional[int] = None,
) -> List[dict]:
    q = db.query(Project).filter(Project.project_type == "smart_contract")
    if user.role == "owner":
        q = q.filter(or_(Project.user_id == user.id, Project.created_by_id == user.id))
    if statuses:
        q = q.filter(Project.status.in_(statuses))
    if owner_id:
        q = q.filter(Project.user_id == owner_id)
    rows = []
    for p in q.order_by(Project.created_at.desc()).all():
        kind = sc_type_of(p)
        if sc_type in ("smart", "leadership") and kind != sc_type:
            continue
        created = _date(p.created_at)
        if date_from and created and created < date_from:
            continue
        if date_to and created and created > date_to:
            continue
        scd = p.smart_contract_data or {}
        ms = scd.get("milestones") or []
        if kind == "leadership":
            reward = _num((scd.get("leadershipTotals") or {}).get("grand")) or sum(_num(m.get("rewardRub")) for m in ms)
            coins = None
        else:
            reward = sum(_num(m.get("rewardRub")) for m in ms)
            coins = sum(_num(m.get("coins")) for m in ms)
        active = route_service.active_step(db, p.id, "contract")
        if active is not None:
            from .approval_roles import ROLE_LABELS
            stage = f"{ROLE_LABELS.get(active.role_key, active.role_key)}: " \
                    f"{active.assignee.full_name if active.assignee else '—'}"
        else:
            stage = {"approved": "Согласован всеми", "rework_needed": "Возвращён на доработку"}.get(p.status or "", "—")
        done = sum(1 for m in ms if m.get("status") in ("paid", "achieved"))
        rows.append({
            "id": p.id,
            "name": p.name or "(без названия)",
            "type": SC_TYPE_LABELS[kind],
            "sc_type": kind,
            "status": STATUS_LABELS.get(p.status or "draft", p.status or ""),
            "stage": stage,
            "owner": p.owner or "",
            "customer": scd.get("customerName") or "",
            "curator": scd.get("curator") or "",
            "business_unit": p.business_unit or "",
            "reward_rub": round(reward, 2),
            "coins": round(coins) if coins is not None else None,
            "rate": _num(scd.get("navRate")) if kind == "smart" and scd.get("navRate") else None,
            "milestones": f"{done}/{len(ms)}",
            "created": _fmt_date(p.created_at),
            "deadline": _fmt_date(scd.get("deadline")),
        })
    for i, r in enumerate(rows, start=1):
        r["num"] = i
    return rows


def describe_filters(sc_type, statuses, date_from, date_to, owner_name) -> str:
    parts = []
    parts.append({"smart": "смарт-контракты", "leadership": "лидерские контракты"}.get(sc_type or "", "все типы"))
    if statuses:
        parts.append("статусы: " + ", ".join(STATUS_LABELS.get(s, s) for s in statuses))
    if date_from or date_to:
        parts.append(f"создан: {_fmt_date(date_from) or '…'} – {_fmt_date(date_to) or '…'}")
    if owner_name:
        parts.append(f"ответственный: {owner_name}")
    return "; ".join(parts)


def build_excel(rows: List[dict], filters_text: str) -> io.BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from .export_excel import _BORDER, _apply_header, _auto_width

    wb = Workbook()
    ws = wb.active
    ws.title = "Реестр контрактов"
    ws.cell(row=1, column=1, value="Реестр контрактов").font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value=f"Сформирован {dt.datetime.now():%d.%m.%Y %H:%M} · {filters_text}").font = Font(italic=True, color="666666")
    _apply_header(ws, [title for _, title in COLUMNS], row=4)
    ws.freeze_panes = "A5"
    for r in rows:
        ws.append([r.get(key) for key, _ in COLUMNS])
        for col_idx in range(1, len(COLUMNS) + 1):
            cell = ws.cell(row=ws.max_row, column=col_idx)
            cell.border = _BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    money_cols = [i for i, (k, _) in enumerate(COLUMNS, start=1) if k in ("reward_rub", "coins", "rate")]
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row):
        for idx in money_cols:
            row[idx - 1].number_format = "#,##0.00" if COLUMNS[idx - 1][0] == "rate" else "#,##0"
    total_row = ws.max_row + 1
    ws.cell(row=total_row, column=2, value=f"Итого: {len(rows)} контракт(ов)").font = Font(bold=True)
    reward_col = [i for i, (k, _) in enumerate(COLUMNS, start=1) if k == "reward_rub"][0]
    coins_col = reward_col + 1
    ws.cell(row=total_row, column=reward_col, value=sum(r["reward_rub"] or 0 for r in rows)).font = Font(bold=True)
    ws.cell(row=total_row, column=reward_col).number_format = "#,##0"
    ws.cell(row=total_row, column=coins_col, value=sum(r["coins"] or 0 for r in rows)).font = Font(bold=True)
    ws.cell(row=total_row, column=coins_col).number_format = "#,##0"
    _auto_width(ws, max_width=40)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


_PDF_CSS = """
@page { size: A4 landscape; margin: 12mm 10mm 14mm 10mm;
  @bottom-center { content: "Страница " counter(page) " из " counter(pages); font-size: 8pt; color: #888; } }
* { box-sizing: border-box; }
body { font-family: 'Plus Jakarta Sans', Arial, sans-serif; font-size: 8pt; color: #1C1C1E; }
h1 { font-size: 16pt; color: #5E5CE6; margin: 0 0 2mm; }
.meta { color: #666; font-size: 8.5pt; margin-bottom: 5mm; }
.kpis { display: flex; gap: 4mm; margin-bottom: 5mm; }
.kpi { background: #F4F3FF; border-radius: 3mm; padding: 3mm 4mm; min-width: 40mm; }
.kpi b { display: block; font-size: 12pt; color: #5E5CE6; }
table { width: 100%; border-collapse: collapse; }
th { background: #5E5CE6; color: #fff; text-align: left; padding: 2mm; font-size: 7.5pt; }
td { padding: 1.8mm 2mm; border-bottom: 0.3mm solid #E5E5ED; vertical-align: top; }
tr:nth-child(even) td { background: #FAFAFC; }
td.num { text-align: right; white-space: nowrap; }
.t-lead { color: #0F766E; font-weight: 700; } .t-smart { color: #5E5CE6; font-weight: 700; }
"""


def build_pdf(rows: List[dict], filters_text: str, generated_by: str) -> io.BytesIO:
    from weasyprint import HTML

    def esc(v):
        return html.escape("" if v is None else str(v))

    def money(v):
        return "" if v is None else f"{v:,.0f}".replace(",", " ")

    total_rub = sum(r["reward_rub"] or 0 for r in rows)
    total_coins = sum(r["coins"] or 0 for r in rows)
    body = [
        "<h1>Реестр контрактов</h1>",
        f"<div class='meta'>Сформирован {dt.datetime.now():%d.%m.%Y %H:%M} · {esc(generated_by)} · {esc(filters_text)}</div>",
        "<div class='kpis'>",
        f"<div class='kpi'>Контрактов<b>{len(rows)}</b></div>",
        f"<div class='kpi'>Вознаграждение<b>{money(total_rub)} ₽</b></div>",
        f"<div class='kpi'>Coin<b>{money(total_coins)}</b></div>",
        f"<div class='kpi'>На согласовании<b>{sum(1 for r in rows if r['status'] == STATUS_LABELS['pending_approval'])}</b></div>",
        "</div>",
        "<table><thead><tr>",
    ]
    cols = [c for c in COLUMNS if c[0] != "business_unit"]
    body += [f"<th>{esc(t)}</th>" for _, t in cols]
    body.append("</tr></thead><tbody>")
    for r in rows:
        body.append("<tr>")
        for key, _ in cols:
            v = r.get(key)
            if key in ("reward_rub", "coins"):
                body.append(f"<td class='num'>{money(v)}</td>")
            elif key == "rate":
                body.append(f"<td class='num'>{'' if v is None else esc(v)}</td>")
            elif key == "type":
                body.append(f"<td class='{'t-lead' if r['sc_type'] == 'leadership' else 't-smart'}'>{esc(v)}</td>")
            else:
                body.append(f"<td>{esc(v)}</td>")
        body.append("</tr>")
    if not rows:
        body.append(f"<tr><td colspan='{len(cols)}' style='text-align:center;color:#888'>Нет контрактов по выбранным фильтрам</td></tr>")
    body.append("</tbody></table>")
    doc = f"<html><head><meta charset='utf-8'><style>{_PDF_CSS}</style></head><body>{''.join(body)}</body></html>"
    return io.BytesIO(HTML(string=doc).write_pdf())
